#!/usr/bin/env python3
r"""
Score the human spot check of pattern precision (data/pattern_precision/
spot_check.xlsx, filled in by hand) against the LLM judgments: agreement and
Cohen's kappa per pattern, and precision per pattern and corpus by the human
and by the LLM on the same items.

Answers come from the spot-check web page (spot_check_answers.jsonl, one
{"n", "answer", "note"} per item) or from the spreadsheet (.xlsx).

Usage:
    python3 spot_check_agreement.py ../data/pattern_precision/spot_check_answers.jsonl \
        ../data/pattern_precision/spot_check_key.jsonl
"""
import argparse
import collections
import json

from openpyxl import load_workbook


def kappa(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    pa, pb = sum(a) / n, sum(b) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("key")
    args = ap.parse_args()
    key = {k["n"]: k for k in map(json.loads, open(args.key))}
    if args.xlsx.endswith(".xlsx"):
        ws = load_workbook(args.xlsx)["Spot check"]
        answers = [(n, answer) for n, _, answer, *_ in ws.iter_rows(min_row=2, values_only=True)]
    else:
        answers = [(r["n"], r["answer"]) for r in map(json.loads, open(args.xlsx))]
    rows = []
    for n, answer in answers:
        a = (answer or "").strip().lower()
        if a in ("yes", "no"):
            rows.append({**key[n], "human": a})
    print(f"{len(rows)}/{len(key)} rows answered")
    for pid in sorted({r["pattern_id"] for r in rows}):
        rs = [r for r in rows if r["pattern_id"] == pid]
        h = [r["human"] == "yes" for r in rs]
        m = [r["llm"] == "yes" for r in rs]
        print(f"\n{pid}: agreement {sum(x == y for x, y in zip(h, m))}/{len(rs)}, kappa {kappa(h, m):.2f}")
        for corpus in ("acl2019", "arxiv2026"):
            c = [r for r in rs if r["corpus"] == corpus]
            if c:
                print(f"  {corpus}: precision human {sum(r['human'] == 'yes' for r in c)}/{len(c)}, "
                      f"LLM {sum(r['llm'] == 'yes' for r in c)}/{len(c)}")
        dis = collections.Counter((r["llm"], r["human"]) for r in rs if r["llm"] != r["human"])
        print("  disagreements (LLM, human):", dict(dis))


if __name__ == "__main__":
    main()
